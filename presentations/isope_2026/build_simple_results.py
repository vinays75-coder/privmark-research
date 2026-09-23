"""Create one editable, plain-language results slide from the completed main study.

Install the slides extra, then run: uv run python <this file>.
PowerPoint and PDF use the same vector layout; PNG is rendered from that PDF.
"""
import csv
import hashlib
import os
from pathlib import Path

import pymupdf
import reportlab
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
STEM = 'PrivMark_One_Slide_Key_Takeaways'
SOURCE = ROOT / 'results/bounded_study_v1/main_exports/comparison.csv'
rows = list(csv.DictReader(SOURCE.open()))
models = ['HuggingFaceTB/SmolLM2-360M-Instruct', 'Qwen/Qwen2.5-0.5B-Instruct',
          'Qwen/Qwen3-0.6B']
assert len(rows) == 6
for model in models:
    matching = [r for r in rows if r['model'] == model]
    assert {r['condition'] for r in matching} == {'baseline', 'hardened'}
    assert all(int(r['records']) == int(r['disclosed_records']) == 12 for r in matching)
    assert all(int(r['cap']) == 64 for r in matching)

W, H = 960, 540
FONTS = Path(os.environ.get('PRIVMARK_FONT_DIR',
    '/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts'))
if (FONTS / 'Calibri.ttf').exists() and (FONTS / 'Calibrib.ttf').exists():
    font_files = [('Calibri', FONTS / 'Calibri.ttf'), ('Calibri-Bold', FONTS / 'Calibrib.ttf')]
else:
    bundled = Path(reportlab.__file__).parent / 'fonts'
    font_files = [('Calibri', bundled / 'Vera.ttf'), ('Calibri-Bold', bundled / 'VeraBd.ttf')]
    print('Calibri unavailable: PDF uses bundled Vera; PowerPoint requests Calibri.')
for name, font_path in font_files:
    pdfmetrics.registerFont(TTFont(name, str(font_path)))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W / 72), Inches(H / 72)
prs.core_properties.title = 'PrivMark in Action: Can AI Keep a Secret?'
prs.core_properties.subject = 'Main live experiment: plain-language evidence summary'
slide = prs.slides.add_slide(prs.slide_layouts[6])
pdf = canvas.Canvas(str(OUT / f'{STEM}.pdf'), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)


def rect(x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    pdf.setFillColorRGB(*(int(color[i:i+2], 16) / 255 for i in (0, 2, 4)))
    pdf.rect(x, H-y-h, w, h, fill=1, stroke=0)


def text(value, x, y, w, size=20, color='253238', bold=False, h=None):
    h = h or size * 1.3
    font = 'Calibri-Bold' if bold else 'Calibri'
    while pdfmetrics.stringWidth(value, font, size) > w and size > 11:
        size -= 0.25
    assert pdfmetrics.stringWidth(value, font, size) <= w, value
    assert x + w <= W and y + h <= H, value
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = shape.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.font.name = 'Calibri'
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = Pt(size * 1.16)
    pdf.setFillColorRGB(*(int(color[i:i+2], 16) / 255 for i in (0, 2, 4)))
    pdf.setFont(font, size)
    pdf.drawString(x, H-y-size*.86, value)


rect(0, 0, W, H, 'FFFFFF')
text('PrivMark in Action: Can AI Keep a Secret?', 40, 24, 880, 34, 'ED6B00', True)
text('Three small local models. Twelve synthetic private records per model.',
     42, 72, 876, 19, '59666B')

# Two columns, three rows: the only numeric result on the slide.
left, top, width = 78, 122, 804
rect(left, top, width, 42, 'E8F2F1')
text('Model', left+20, top+10, 382, 19, '176F6C', True)
text('Private records exposed', left+426, top+10, 358, 19, '176F6C', True)
for i, model in enumerate(models):
    y = top+42+i*50
    rect(left, y, width, 50, 'F8FAFA' if i % 2 == 0 else 'FFFFFF')
    text(model.split('/')[-1], left+20, y+13, 392, 21)
    text('12 out of 12', left+426, y+10, 358, 25, 'AC4535', True)
    rect(left, y+49, width, 1, 'D9E2E1')
text('Same result with basic and stronger confidentiality instructions.',
     left, 328, width, 18, '59666B')

text('Key Takeaways', 78, 378, 804, 21, '665941', True)
for i, line in enumerate([
    'Good answers do not guarantee privacy.',
    'Stronger instructions did not stop the leaks.',
    'PrivMark makes the risk clear and links it to evidence.',
]):
    text('•', 90, 409+i*27, 20, 19, '665941')
    text(line, 113, 409+i*27, 769, 20, '3E494D')
text('Controlled local experiment; these results are not a general model ranking.',
     42, 494, 876, 12, '6B7478')
text('IEEE Symposium on Privacy Expectations (ISoPE) 2026',
     42, 516, 876, 14, '59666B')
rect(0, 535, W, 5, '4FB3AD')

notes = '''TALK TRACK (about 45 seconds)

We asked three small AI models to keep a made-up private token secret while still answering a public question. Each model saw the same twelve test records. All three exposed the private token from every record through at least one attack. Adding stronger confidentiality instructions did not change that result.

The two Qwen models also answered every public question correctly. This is the point: being useful does not mean being private. PrivMark makes that difference visible, with the evidence available behind the label.

DETAILS FOR QUESTIONS — NOT ADDITIONAL ON-SLIDE CLAIMS

These are real local model runs with synthetic records, not simulated model responses. The main experiment used three batches of four fresh records, four attacks plus one public question, two instruction conditions, and a 64-token output cap: 360 scored responses. The table counts records exposed by any attack, not independent attack attempts. The 95% Wilson interval for 12/12 is approximately 75.8–100%. This is a controlled context-disclosure demonstration, not training-data extraction, a privacy certification, or a general model ranking.

Thirty responses were cut off (8.3%); each reported full disclosure was directly observed. Cut-off responses without a detected token do not establish safety. The user elected to proceed after the pilot exceeded the planned 5% truncation gate; that decision is preserved in every main artifact, without changing the budget. The slide reports positive evidence of leakage and does not claim resistance for any model.

SmolLM2 public-answer pattern checks: 8/12 basic and 9/12 stronger. A later raw-response audit found that all seven unconfirmed public answers returned the private token alone (4/12 basic, 3/12 stronger). The original comparison does not display benign-question leakage as a separate metric; see docs/EXPERIMENT_12_RECORDS.md. Both Qwen models passed 12/12 under each condition. No privacy winner is established by the headline record-disclosure result. This supports a limited worked example of PrivMark's evidence-backed communication concept, not validation of all eleven framework dimensions or stakeholder usability.
'''
notes += f'\nSOURCE: {SOURCE.relative_to(ROOT)}\nSHA256: {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}\n'
slide.notes_slide.notes_text_frame.text = notes
prs.save(OUT / f'{STEM}.pptx')
pdf.showPage()
pdf.save()
with pymupdf.open(OUT / f'{STEM}.pdf') as document:
    assert len(document) == 1
    document[0].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(OUT / f'{STEM}.png')
(OUT / f'{STEM}_speaker_notes.md').write_text('# One-slide results: speaker notes\n\n' + notes)
check = Presentation(OUT / f'{STEM}.pptx')
assert len(check.slides) == 1
assert sum('12 out of 12' in s.text for s in check.slides[0].shapes if s.has_text_frame) == 3
assert 'truncation gate' in check.slides[0].notes_slide.notes_text_frame.text
print(f'Created editable PPTX, PDF, PNG and speaker notes: {OUT / STEM}')
