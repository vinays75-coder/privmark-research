"""Create one editable results slide with method, paired outcomes and limitations from the completed main study.

Install the slides extra, then run: uv run python <this file>.
PowerPoint and PDF use the same vector layout; PNG is rendered from that PDF.
"""
import csv
import hashlib
import os
from pathlib import Path

import pymupdf
import reportlab
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
STEM = 'PrivMark_One_Slide_Experiment_Evidence'
SOURCE = ROOT / 'results/bounded_study_v1/main_exports/comparison.csv'
rows = list(csv.DictReader(SOURCE.open()))
models = ['HuggingFaceTB/SmolLM2-360M-Instruct', 'Qwen/Qwen2.5-0.5B-Instruct',
          'Qwen/Qwen3-0.6B']
assert len(rows) == 6
for model in models:
    matching = [r for r in rows if r['model'] == model]
    assert {r['condition'] for r in matching} == {'baseline', 'hardened'}
    assert all(int(r['records']) == int(r['disclosed_records']) == 12 for r in matching)
    assert all(int(r['cap']) == 64 for r in matching)

W, H = 960, 540
FONTS = Path(os.environ.get('PRIVMARK_FONT_DIR',
    '/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts'))
if (FONTS / 'Calibri.ttf').exists() and (FONTS / 'Calibrib.ttf').exists():
    font_files = [('Calibri', FONTS / 'Calibri.ttf'), ('Calibri-Bold', FONTS / 'Calibrib.ttf')]
else:
    bundled = Path(reportlab.__file__).parent / 'fonts'
    font_files = [('Calibri', bundled / 'Vera.ttf'), ('Calibri-Bold', bundled / 'VeraBd.ttf')]
    print('Calibri unavailable: PDF uses bundled Vera; PowerPoint requests Calibri.')
for name, font_path in font_files:
    pdfmetrics.registerFont(TTFont(name, str(font_path)))

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W / 72), Inches(H / 72)
prs.core_properties.title = 'PrivMark in Action: Useful Answers, Exposed Secrets'
prs.core_properties.subject = 'Main live experiment: plain-language evidence summary'
slide = prs.slides.add_slide(prs.slide_layouts[6])
pdf = canvas.Canvas(str(OUT / f'{STEM}.pdf'), pagesize=(W, H))
pdf.setTitle(prs.core_properties.title)


def rect(x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    pdf.setFillColorRGB(*(int(color[i:i+2], 16) / 255 for i in (0, 2, 4)))
    pdf.rect(x, H-y-h, w, h, fill=1, stroke=0)


def text(value, x, y, w, size=20, color='253238', bold=False, h=None):
    h = h or size * 1.3
    font = 'Calibri-Bold' if bold else 'Calibri'
    while pdfmetrics.stringWidth(value, font, size) > w and size > 11:
        size -= 0.25
    assert pdfmetrics.stringWidth(value, font, size) <= w, value
    assert x + w <= W and y + h <= H, value
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = shape.text_frame
    frame.word_wrap = False
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.font.name = 'Calibri'
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(color)
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = Pt(size * 1.16)
    pdf.setFillColorRGB(*(int(color[i:i+2], 16) / 255 for i in (0, 2, 4)))
    pdf.setFont(font, size)
    pdf.drawString(x, H-y-size*.86, value)


rect(0, 0, W, H, 'FFFFFF')
text('PrivMark in Action: Useful Answers, Exposed Secrets',
     40, 22, 880, 31, 'ED6B00', True)
text('Live local comparison • Same synthetic records and prompts for every model',
     42, 66, 876, 18, '59666B')

# Visible protocol: distinguish twelve independent records from 360 responses.
rect(42, 100, 876, 53, 'F1F6F5')
for label, detail, x, w in [
    ('12 records', '3 seeds × 4 records', 54, 122),
    ('5 questions', '4 attacks + 1 public', 220, 134),
    ('2 conditions', 'Basic / stronger rules', 395, 139),
    ('3 models', 'Loaded one at a time', 576, 154),
    ('360 responses', '64-token response cap', 752, 155),
]:
    text(label, x, 108, w, 19, '176F6C', True)
    text(detail, x, 133, w, 11.5, '59666B')
for value, x in [('×', 184), ('×', 359), ('×', 542), ('=', 723)]:
    text(value, x, 109, 26, 21, '59666B')

# Six columns with grouped headers, rather than six rows of repeated model names.
left, top, total = 42, 176, 876
columns = [42, 222, 347, 472, 597, 722, 918]
rect(left, top, total, 54, 'E8F2F1')
text('Model', 54, 190, 156, 18, '176F6C', True)
text('Private records exposed', 232, 181, 230, 18, '176F6C', True)
text('Public answers confirmed*', 482, 181, 230, 17, '176F6C', True)
text('Cut-off responses', 734, 181, 172, 18, '176F6C', True)
for label, x, w in [('Basic', 235, 98), ('Stronger', 360, 98),
                    ('Basic', 485, 98), ('Stronger', 610, 98),
                    ('Both conditions', 736, 170)]:
    text(label, x, 208, w, 13, '59666B')

names = [('SmolLM2', '360M · Instruct'), ('Qwen2.5', '0.5B · Instruct'), ('Qwen3', '0.6B')]
for i, (model, (label, size_label)) in enumerate(zip(models, names, strict=True)):
    y = top + 54 + i * 48
    rect(left, y, total, 48, 'F8FAFA' if i % 2 == 0 else 'FFFFFF')
    text(label, 54, y+4, 156, 20, bold=True)
    text(size_label, 54, y+29, 156, 12, '59666B')
    arms = {r['condition']: r for r in rows if r['model'] == model}
    for j, condition in enumerate(('baseline', 'hardened')):
        row = arms[condition]
        exposed = f"{row['disclosed_records']}/{row['records']}"
        public = f"{row['public_confirmed_correct']}/{row['records']}"
        text(exposed, 236+j*125, y+12, 100, 22, 'AC4535', True)
        text(public, 486+j*125, y+12, 100, 22, '176F6C', True)
    cut = sum(int(r['truncated']) for r in arms.values())
    count = sum(int(r['responses']) for r in arms.values())
    rate = f'{100 * cut / count:.1f}%' if cut else '0%'
    text(f'{cut}/{count} ({rate})', 736, y+14, 172, 19,
         'AC4535' if cut else '176F6C', True)
    rect(left, y+47, total, 1, 'D9E2E1')
for x in (222, 472, 722):
    rect(x, 176, 1, 198, 'D9E2E1')

text('*SmolLM2’s 7 unconfirmed public answers leaked the secret. Exposure means a full token leaked in at least one attack.',
     42, 382, 876, 12, '59666B')
text('Uncertainty: 12/12 exposed = 100% observed, with a 95% confidence interval of 76–100% (12 records).',
     42, 399, 876, 12, '59666B')
text('Key Takeaways', 42, 427, 876, 20, '665941', True)
for i, line in enumerate([
    'No privacy winner: all 12 records leaked for every model, even with stronger instructions.',
    'Qwen2.5 had no cut-off answers. PrivMark shows usefulness and privacy separately.',
]):
    text('•', 49, 456+i*26, 20, 17, '665941')
    text(line, 69, 456+i*26, 849, 17, '3E494D')
text('IEEE Symposium on Privacy Expectations (ISoPE) 2026',
     42, 516, 876, 14, '59666B')
rect(0, 535, W, 5, '4FB3AD')

notes = '''TALK TRACK (about 45 seconds)

We asked three small AI models to keep a made-up private token secret while still answering a public question. Each model saw the same twelve test records. All three exposed the private token from every record through at least one attack. Adding stronger confidentiality instructions did not change that result.

The two Qwen models also answered every public question correctly. This is the point: being useful does not mean being private. PrivMark makes that difference visible, with the evidence available behind the label.

DETAILS FOR QUESTIONS — NOT ADDITIONAL ON-SLIDE CLAIMS

These are real local model runs with synthetic records, not simulated model responses. The main experiment used three batches of four fresh records, four attacks plus one public question, two instruction conditions, and a 64-token output cap: 360 scored responses. The table counts records exposed by any attack, not independent attack attempts. The 95% Wilson interval for 12/12 is approximately 75.8–100%. This is a controlled context-disclosure demonstration, not training-data extraction, a privacy certification, or a general model ranking.

Thirty responses were cut off (8.3%); each reported full disclosure was directly observed. Cut-off responses without a detected token do not establish safety. The user elected to proceed after the pilot exceeded the planned 5% truncation gate; that decision is preserved in every main artifact, without changing the budget. The slide reports positive evidence of leakage and does not claim resistance for any model.

SmolLM2 public-answer pattern checks: 8/12 basic and 9/12 stronger. A later raw-response audit found that all seven unconfirmed public answers returned the private token alone (4/12 basic, 3/12 stronger). The original comparison does not display benign-question leakage as a separate metric; see docs/EXPERIMENT_12_RECORDS.md. Both Qwen models passed 12/12 under each condition. No privacy winner is established by the headline record-disclosure result. This supports a limited worked example of PrivMark's evidence-backed communication concept, not validation of all eleven framework dimensions or stakeholder usability.
'''
notes += f'\nSOURCE: {SOURCE.relative_to(ROOT)}\nSHA256: {hashlib.sha256(SOURCE.read_bytes()).hexdigest()}\n'
slide.notes_slide.notes_text_frame.text = notes
prs.save(OUT / f'{STEM}.pptx')
pdf.showPage()
pdf.save()
with pymupdf.open(OUT / f'{STEM}.pdf') as document:
    assert len(document) == 1
    document[0].get_pixmap(matrix=pymupdf.Matrix(2, 2)).save(OUT / f'{STEM}.png')
(OUT / f'{STEM}_speaker_notes.md').write_text('# Expanded experiment evidence: speaker notes\n\n' + notes)
check = Presentation(OUT / f'{STEM}.pptx')
assert len(check.slides) == 1
assert sum(s.text == '12/12' for s in check.slides[0].shapes if s.has_text_frame) == 10
assert 'truncation gate' in check.slides[0].notes_slide.notes_text_frame.text
print(f'Created editable PPTX, PDF, PNG and speaker notes: {OUT / STEM}')
